"""Stage 3 - read the corpus, chunk it, embed it, store it locally.

Run:  python src/ingest.py
"""
import os
from docx import Document
from sentence_transformers import SentenceTransformer
import chromadb

CORPUS_DIR = "corpus"

# Sized for THIS corpus, deliberately. The policy documents run 265-400
# words each, so a 500-word chunk would swallow every document whole and
# the overlap below would never do anything. At 200/40 each document
# splits into 2-3 overlapping chunks, which is what makes retrieval
# operate on passages rather than whole files.
#
# The general lesson: chunk size is a property of your corpus, not a
# universal default. Print len(chunk(text)) after your first run - if
# every document yields exactly one chunk, your chunking is inert.
CHUNK_WORDS = 200
OVERLAP_WORDS = 40

model = SentenceTransformer("all-MiniLM-L6-v2")
client = chromadb.PersistentClient(path="vectordb")
collection = client.get_or_create_collection("policies")


def read_docx(path):
    """Extract paragraph text and table text from a Word file."""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def chunk(text):
    """Overlapping windows. Overlap stops a sentence being orphaned
    across a boundary where neither half makes sense alone."""
    words = text.split()
    step = CHUNK_WORDS - OVERLAP_WORDS
    return [
        " ".join(words[i:i + CHUNK_WORDS])
        for i in range(0, len(words), step)
        if words[i:i + CHUNK_WORDS]
    ]


def main():
    if not os.path.isdir(CORPUS_DIR):
        raise SystemExit(f"No '{CORPUS_DIR}' folder. Download the Policies library into it first.")

    doc_id = 0
    for filename in sorted(os.listdir(CORPUS_DIR)):
        if not filename.endswith(".docx") or filename.startswith("~$"):
            continue
        text = read_docx(os.path.join(CORPUS_DIR, filename))
        pieces = chunk(text)
        for piece in pieces:
            collection.add(
                ids=[f"chunk-{doc_id}"],
                documents=[piece],
                embeddings=[model.encode(piece).tolist()],
                metadatas=[{"source": filename}],
            )
            doc_id += 1
        print(f"Indexed {filename} ({len(pieces)} chunks)")

    print(f"\nDone. {doc_id} chunks stored in ./vectordb")


if __name__ == "__main__":
    main()
